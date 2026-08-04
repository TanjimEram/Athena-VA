"""Writing documents for real - Google Docs through the API, with a local
Word file as the fallback when there's no internet or no Google sign-in.

These functions only move text in and out of documents. The PROSE always
comes from the brain: it composes what to say, then hands the finished text
to write_to_document. Nothing here writes content of its own.

A document can be named by its Google id or by its title - anything that
isn't an id is looked up by name first. Because Athena holds the drive.file
scope, she can only see documents SHE created, never the rest of your Drive.

After a successful write the document opens in the browser so you can watch
it appear. Every function returns one short honest sentence, spoken verbatim
- a failure is always reported as a failure."""

import os
import re
import webbrowser

from athena import google_auth

DOC_URL = "https://docs.google.com/document/d/{}/edit"
# A Google file id: long, no spaces. Titles never look like this, which is
# how we tell "write to 1a2B3c..." from "write to my shopping list".
_ID_RE = re.compile(r"^[A-Za-z0-9_-]{25,}$")
# read_document is spoken aloud, so a long document has to be cut short.
MAX_SPOKEN_CHARS = 1200
# The test script turns this off so a run doesn't open a pile of tabs.
OPEN_IN_BROWSER = True

# The last document touched, so "add a line to it" works without repeating
# the title. Only used when the caller gives no usable identifier.
_last_doc_id = None
_last_doc_title = ""

_PRONOUNS = {"it", "that", "this", "the document", "the doc", "same",
             "that one", "it again", ""}


def _remember(doc_id: str, title: str) -> None:
    global _last_doc_id, _last_doc_title
    _last_doc_id, _last_doc_title = doc_id, title


def _escape(text: str) -> str:
    """Drive's query language is single-quoted, so quotes need escaping."""
    return text.replace("\\", "\\\\").replace("'", "\\'")


def _search(title: str) -> list:
    """Documents whose name matches `title`. [] if none or not signed in."""
    drive = google_auth.get_service("drive", "v3")
    if drive is None:
        return []
    query = (f"name contains '{_escape(title)}' and "
             "mimeType='application/vnd.google-apps.document' and "
             "trashed=false")
    try:
        found = drive.files().list(
            q=query, fields="files(id,name,modifiedTime)",
            orderBy="modifiedTime desc", pageSize=10).execute()
        return found.get("files", [])
    except Exception as exc:
        print(f"[documents] search failed: {google_auth._short(exc)}")
        return []


def _resolve(doc_id_or_title: str):
    """Turn whatever the brain said into (doc_id, title, error_sentence).
    Exactly one of doc_id / error_sentence is set."""
    raw = (doc_id_or_title or "").strip()

    if raw.lower().strip(" .") in _PRONOUNS:
        if _last_doc_id:
            return _last_doc_id, _last_doc_title, None
        return None, "", "I'm not sure which document you mean."

    if _ID_RE.match(raw):
        return raw, "", None

    matches = _search(raw)
    if not matches:
        return None, "", (f"I couldn't find a document called {raw}. I can "
                          "only see documents I created myself.")
    if len(matches) > 1:
        names = ", ".join(m["name"] for m in matches[:3])
        exact = [m for m in matches if m["name"].lower() == raw.lower()]
        if len(exact) != 1:
            return None, "", (f"I found more than one document matching {raw}: "
                              f"{names}. Which one did you mean?")
        matches = exact
    return matches[0]["id"], matches[0]["name"], None


def _doc_end_index(docs, doc_id: str) -> int:
    """Where the body currently ends. 2 means the document is empty."""
    doc = docs.documents().get(documentId=doc_id).execute()
    content = doc.get("body", {}).get("content", [])
    return content[-1].get("endIndex", 2) if content else 2


def _open(doc_id: str) -> None:
    if OPEN_IN_BROWSER:
        try:
            webbrowser.open(DOC_URL.format(doc_id))
        except Exception as exc:
            print(f"[documents] couldn't open the browser: {exc}")


def create_document(title: str) -> str:
    """Make a new, empty Google Doc called `title`."""
    docs = google_auth.get_service("docs", "v1")
    if docs is None:
        return google_auth.not_connected_message()
    title = (title or "Untitled document").strip()
    try:
        doc = docs.documents().create(body={"title": title}).execute()
    except Exception as exc:
        return f"I couldn't create that document: {google_auth._short(exc)}"
    doc_id = doc.get("documentId", "")
    _remember(doc_id, title)
    print(f"[documents] created {title!r} ({doc_id})")
    return f"I've created a document called {title}."


def write_to_document(doc_id_or_title: str, text: str, mode: str = "append") -> str:
    """Put `text` into a document. mode 'append' adds to the end, 'replace'
    clears it first. Opens the document in the browser afterwards."""
    docs = google_auth.get_service("docs", "v1")
    if docs is None:
        return google_auth.not_connected_message()
    if not (text or "").strip():
        return "There was nothing for me to write, so I left the document alone."

    doc_id, title, error = _resolve(doc_id_or_title)
    if error:
        return error

    mode = "replace" if str(mode).lower().startswith("replace") else "append"
    try:
        end = _doc_end_index(docs, doc_id)
        requests = []
        if mode == "replace" and end > 2:
            requests.append({"deleteContentRange": {
                "range": {"startIndex": 1, "endIndex": end - 1}}})
            body = text
        else:
            # Start appended text on its own line, unless the doc is empty.
            body = ("\n" + text) if (mode == "append" and end > 2) else text
        requests.append({"insertText": {"endOfSegmentLocation": {},
                                        "text": body}})
        docs.documents().batchUpdate(
            documentId=doc_id, body={"requests": requests}).execute()
    except Exception as exc:
        return f"I couldn't write to that document: {google_auth._short(exc)}"

    _remember(doc_id, title)
    _open(doc_id)
    where = f" to {title}" if title else ""
    words = len(text.split())
    if mode == "replace":
        return f"I've replaced the contents{where} with {words} words, and opened it."
    return f"I've added {words} words{where}, and opened it so you can see."


def find_document(title: str) -> str:
    """Look for documents by name and say what turned up."""
    if google_auth.get_service("drive", "v3") is None:
        return google_auth.not_connected_message()
    title = (title or "").strip()
    if not title:
        return "Tell me what the document is called and I'll look for it."

    matches = _search(title)
    if not matches:
        return (f"I couldn't find a document called {title}. I can only see "
                "documents I created myself.")
    if len(matches) == 1:
        _remember(matches[0]["id"], matches[0]["name"])
        return f"I found one document: {matches[0]['name']}."
    names = ", ".join(m["name"] for m in matches[:5])
    more = " and a few others" if len(matches) > 5 else ""
    return f"I found {len(matches)} documents: {names}{more}."


def read_document(doc_id_or_title: str) -> str:
    """Read a document back. Long documents are cut short - this is spoken."""
    docs = google_auth.get_service("docs", "v1")
    if docs is None:
        return google_auth.not_connected_message()

    doc_id, title, error = _resolve(doc_id_or_title)
    if error:
        return error

    try:
        doc = docs.documents().get(documentId=doc_id).execute()
    except Exception as exc:
        return f"I couldn't read that document: {google_auth._short(exc)}"

    title = doc.get("title", title)
    text = _extract_text(doc)
    _remember(doc_id, title)
    if not text.strip():
        return f"{title} is empty."
    if len(text) > MAX_SPOKEN_CHARS:
        return (f"{title} says: {text[:MAX_SPOKEN_CHARS].rstrip()}... "
                "that's as far as I'll read for now.")
    return f"{title} says: {text}"


def _extract_text(doc: dict) -> str:
    """Pull the plain text out of the Docs API's nested structure."""
    parts = []
    for block in doc.get("body", {}).get("content", []):
        paragraph = block.get("paragraph")
        if not paragraph:
            continue                      # tables/images have no spoken text
        for element in paragraph.get("elements", []):
            parts.append(element.get("textRun", {}).get("content", ""))
    return "".join(parts).strip()


def _documents_folder() -> str:
    """The user's Documents folder. OneDrive often moves it, so check there
    too before giving up and using the home folder."""
    home = os.path.expanduser("~")
    for candidate in (os.path.join(home, "Documents"),
                      os.path.join(home, "OneDrive", "Documents")):
        if os.path.isdir(candidate):
            return candidate
    return home


def _safe_filename(title: str) -> str:
    cleaned = re.sub(r'[<>:"/\\|?*]', "", title).strip() or "Untitled"
    return cleaned[:120] + ".docx"


def write_local_docx(title: str, text: str) -> str:
    """Fallback: write a real Word file to the Documents folder and open it.
    Works with no internet and no Google account."""
    if not (text or "").strip():
        return "There was nothing for me to write, so I didn't make the file."
    title = (title or "Untitled document").strip()
    path = os.path.join(_documents_folder(), _safe_filename(title))
    try:
        from docx import Document
        document = Document()
        document.add_heading(title, level=1)
        for paragraph in text.split("\n"):
            document.add_paragraph(paragraph)
        document.save(path)
    except Exception as exc:
        return f"I couldn't save that document: {exc}"

    # Say where it actually went - on a machine with no Documents folder this
    # lands in the home folder, and claiming otherwise would be a lie.
    folder = os.path.basename(os.path.dirname(path))
    where = "your Documents folder" if folder == "Documents" else f"your {folder} folder"
    try:
        os.startfile(path)
    except Exception as exc:
        print(f"[documents] couldn't open {path}: {exc}")
        return f"I've saved {title} to {where}, but couldn't open it."
    print(f"[documents] wrote {path}")
    return f"I've saved {title} to {where} and opened it."
